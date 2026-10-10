# reports/tests.py — Excel/PDF export ও রিমাইন্ডার সহায়কের পরীক্ষা (Django ছাড়াই চলে)
import io
import unittest
from decimal import Decimal

from openpyxl import load_workbook

from . import exports
from .reminder_utils import DEFAULT_TEMPLATE, money, normalize_phone, render_message


class ExportBuilderTests(unittest.TestCase):
    ROWS = [{'invoice_no': 'INV-1', 'sale_date': '2026-10-02T10:00:00+06:00', 'customer_name': 'Rahim', 'sales_by': 'admin',
             'sales_price': '1,000.50', 'cost_price': 800, 'profit': 200.5, 'collect_amount': 500, 'due_amount': 500.5}]

    def test_every_table_config_builds_xlsx_and_pdf_even_when_empty(self):
        for key, cfg in exports.REPORT_EXPORTS.items():
            if cfg.get('kind') == 'statement':
                continue
            cards = [(label, 0, f) for label, _k, f in cfg['cards']]
            self.assertTrue(exports.build_xlsx('Shop', cfg['title'], 'Oct 2026', cfg['cols'], [], cards).startswith(b'PK'), key)
            self.assertTrue(exports.build_pdf('Shop', cfg['title'], 'Oct 2026', cfg['cols'], [], cards).startswith(b'%PDF'), key)

    def test_xlsx_keeps_numbers_numeric_and_dates_real(self):
        cfg = exports.REPORT_EXPORTS['sales']
        data = exports.build_xlsx('Shop', cfg['title'], 'Oct', cfg['cols'], self.ROWS)
        ws = load_workbook(io.BytesIO(data)).active
        rows = list(ws.iter_rows(values_only=True))
        head = next(r for r in rows if r[0] == 'Invoice')
        body = rows[rows.index(head) + 1]
        self.assertEqual(body[0], 'INV-1')
        self.assertEqual(body[1].date().isoformat(), '2026-10-02')
        self.assertEqual(body[4], 1000.5)   # '1,000.50' টেক্সট থেকে সংখ্যা

    def test_sheet_title_with_slash_is_allowed(self):
        stmt = exports.statement_from_data({'total_sales': 10, 'total_purchase': 4, 'gross_profit': 6, 'total_expenses': 1, 'net_profit': 5})
        self.assertEqual(load_workbook(io.BytesIO(exports.build_xlsx('Shop', 'Profit / Loss', 'Oct', [], [], None, stmt))).active.title, 'Profit - Loss')
        self.assertTrue(exports.build_pdf('Shop', 'Profit / Loss', 'Oct', [], [], None, stmt).startswith(b'%PDF'))

    def test_bengali_text_does_not_crash_pdf(self):
        rows = [dict(self.ROWS[0], customer_name='রহিম <b>&')]
        cfg = exports.REPORT_EXPORTS['sales']
        self.assertTrue(exports.build_pdf('মেহেরিন মার্ট', cfg['title'], 'Oct', cfg['cols'], rows).startswith(b'%PDF'))


class ReminderHelperTests(unittest.TestCase):
    def test_phone_normalisation(self):
        for raw in ('01712-345678', '+8801712345678', '8801712345678', '1712345678', '008801712345678', '0171 234 5678'):
            self.assertEqual(normalize_phone(raw), '8801712345678', raw)
        for raw in ('', None, 'N/A', '017123'):
            self.assertIsNone(normalize_phone(raw), raw)

    def test_money_text(self):
        self.assertEqual(money(1250), '1,250')
        self.assertEqual(money(Decimal('1250.50')), '1,250.50')

    def test_message_fills_placeholders_and_ignores_unknown(self):
        out = render_message(DEFAULT_TEMPLATE, name='রহিম', amount=Decimal('1250'), shop='দোকান', phone='017', invoices='', days=3)
        self.assertIn('রহিম', out)
        self.assertIn('৳1,250', out)
        self.assertEqual(render_message('{name} {nope}', name='A', amount=0, shop='', phone='', invoices='', days=0), 'A {nope}')


if __name__ == '__main__':
    unittest.main()
